# -*- coding: utf-8 -*-
"""Excel 读取层（ADR-002）：xlrd 读 .xls、openpyxl 读 .xlsx，统一输出单元格矩阵。

单元格统一表示: {"raw": str(原文), "num": float|None, "date": datetime.date|None}
"""
from __future__ import annotations

import datetime as _dt
import io
from dataclasses import dataclass, field

import xlrd
from openpyxl import load_workbook

EXCEL_SERIAL_MIN = 20000  # 1954 年，早于业务可能的日期（R-IMP-12 序列号启发式）
EXCEL_SERIAL_MAX = 80000  # 2119 年


@dataclass
class SheetMatrix:
    name: str
    rows: list[list[dict]] = field(default_factory=list)

    @property
    def nrows(self) -> int:
        return len(self.rows)

    def ncols(self) -> int:
        return max((len(r) for r in self.rows), default=0)


def _num_to_date(v: float) -> _dt.date | None:
    if EXCEL_SERIAL_MIN <= v <= EXCEL_SERIAL_MAX:
        try:
            base = _dt.datetime(1899, 12, 30)
            return (base + _dt.timedelta(days=v)).date()
        except (OverflowError, ValueError):
            return None
    return None


def _cell(raw, num=None, date=None) -> dict:
    if raw is None:
        raw = ""
    text = str(raw).strip() if not isinstance(raw, str) else raw.strip()
    return {"raw": text, "num": num, "date": date}


def _to_cell(value) -> dict:
    if value is None:
        return _cell("")
    if isinstance(value, _dt.datetime):
        return _cell(value.strftime("%Y-%m-%d %H:%M:%S"), date=value.date())
    if isinstance(value, _dt.date):
        return _cell(value.strftime("%Y-%m-%d"), date=value)
    if isinstance(value, bool):
        return _cell(str(value))
    if isinstance(value, (int, float)):
        d = _num_to_date(float(value))
        return _cell(_fmt_num(float(value)), float(value), d)
    return _cell(str(value))


def _fmt_num(v: float) -> str:
    if v == int(v):
        return str(int(v))
    return str(v)


def parse_date_value(cell: dict) -> _dt.date | None:
    """从单元格提取日期：date 类型 > 序列号启发式 > 文本解析（R-IMP-12）"""
    if cell["date"]:
        return cell["date"]
    if cell["num"] is not None:
        return _num_to_date(cell["num"])
    text = (cell["raw"] or "").strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%m/%d/%Y", "%Y%m%d"):
        try:
            return _dt.datetime.strptime(text.split(" ")[0], fmt).date()
        except ValueError:
            continue
    return None


def read_workbook(filename: str, data: bytes) -> list[SheetMatrix]:
    """按扩展名+签名读取工作簿，返回工作表矩阵列表。失败抛 ValueError。"""
    lower = (filename or "").lower()
    if lower.endswith(".xlsx") or data[:2] == b"PK":
        return _read_xlsx(data)
    if lower.endswith(".xls") or data[:4] == b"\xd0\xcf\x11\xe0":
        return _read_xls(data)
    raise ValueError("不支持的文件格式（仅支持 .xls/.xlsx）")


def _read_xls(data: bytes) -> list[SheetMatrix]:
    try:
        wb = xlrd.open_workbook(file_contents=data)
    except Exception as e:
        raise ValueError(f"xls 文件解析失败（可能加密或损坏）: {e}")
    sheets = []
    for ws in wb.sheets():
        sm = SheetMatrix(name=ws.name)
        for r in range(ws.nrows):
            row = []
            for c in range(ws.ncols):
                cell = ws.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    dt = xlrd.xldate_as_datetime(cell.value, wb.datemode)
                    row.append(_cell(dt.strftime("%Y-%m-%d %H:%M:%S"), date=dt.date()))
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    row.append(_to_cell(cell.value))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(_cell(""))
                else:
                    row.append(_cell(cell.value))
            sm.rows.append(row)
        sheets.append(sm)
    return sheets


def _read_xlsx(data: bytes) -> list[SheetMatrix]:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:
        raise ValueError(f"xlsx 文件解析失败（可能加密或损坏）: {e}")
    sheets = []
    for ws in wb.worksheets:
        sm = SheetMatrix(name=ws.title)
        for row in ws.iter_rows(values_only=True):
            sm.rows.append([_to_cell(v) for v in row])
        sheets.append(sm)
    wb.close()
    return sheets
