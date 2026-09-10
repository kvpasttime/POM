# -*- coding: utf-8 -*-
"""生成手动测试用的 Excel 测试数据（均为虚构，输出到 D:/POM/testdata/）"""
import datetime as dt
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import Workbook

OUT = Path(r"D:\POM\testdata")
OUT.mkdir(exist_ok=True)

TEMPLATE_HEADER = ["行号", "报价行编号", "物料编码", "物料名称", "技术参数", "规格型号", "零件号/图号",
                   "材质", "品牌", "进口/国产", "单位", "使用单位", "需求日期", "备注", "需求数量",
                   "供应商规格型号", "供应商材质", "供应商技术参数", "品牌/厂家",
                   "*交货日期", "*运输方式", "*可供数量", "*含税单价", "报价备注"]


def tr(code, name, tech="", spec="", part="", mat="", brand="", origin="", unit="",
       use_unit="", req_date=None, remark="", qty=None, factory="", transport="",
       avail=None, amount=None, q_remark=""):
    """22 字段（对应模板表头第 3~24 列）"""
    return [code, name, tech, spec, part, mat, brand, "", unit, use_unit,
            req_date, remark, qty, "", "", "", factory, "", transport, avail, amount, q_remark]


def build_template_wb(title_date: str, rows: list) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "导入报价模板"
    ws.append(["导入报价模板"])
    ws.cell(row=2, column=1, value=f"日期:{title_date}")
    ws.append(TEMPLATE_HEADER)
    for i, r in enumerate(rows, start=1):
        assert len(r) == 22, f"行 {i} 字段数 {len(r)} != 22"
        ws.append([i, f"1277000000{i:02d}", *r])
    return wb


# ============ 文件1：标准模板带价格 ============
wb1 = build_template_wb("2026-09-01 10:00:00", [
    tr("TEST-1001", "测试轴承", spec="6205-2RS 深沟球轴承", brand="NSK", unit="套",
       use_unit="测试一矿", req_date=dt.date(2026, 11, 15), qty=10, transport="汽运",
       avail=10, amount=85.5, q_remark="含税含运，现货"),
    tr("TEST-1002", "测试皮带", spec="B型 三角带 2000mm", brand="三星", unit="条",
       use_unit="测试二矿", req_date=dt.date(2026, 12, 1), qty=4, transport="汽运",
       avail=4, amount=38.0, q_remark="含税不含运"),
    tr("TEST-1003", "测试电机", tech="YE2-132M-4 7.5kW", brand="西门子", unit="台",
       use_unit="测试一矿", req_date=dt.date(2026, 11, 20), qty=2, transport="汽运",
       avail=2, amount=4850.0, q_remark="含税含运，货期45天"),
    tr("TEST-1004", "测试电缆", tech="YJV 3x95+1x50", unit="米", use_unit="测试二矿",
       req_date=dt.date(2026, 11, 20), qty=300, transport="汽运", avail=300,
       amount=52.0, q_remark="含税含运"),
    tr("TEST-1005", "测试闸阀", spec="DN50 PN16 Z41H-16C", unit="台", use_unit="测试一矿",
       req_date=dt.date(2026, 12, 10), remark="急用", qty=3, transport="汽运",
       avail=3, amount=620.0, q_remark="含税含运，现货"),
    tr("TEST-1006", "测试滤芯", spec="液压油滤芯 LX-100", brand="黎明", unit="件",
       use_unit="测试一矿", req_date=dt.date(2026, 11, 18), qty=20, transport="汽运",
       avail=20, amount=145.0, q_remark="含税不含运"),
])
wb1.save(OUT / "测试文件1_标准模板_带价格.xlsx")

# ============ 文件2：仅需求无报价 ============
wb2 = build_template_wb("2026-08-25 10:00:00", [
    tr("TEST-2001", "测试链轮", spec="双节距链轮 20齿", unit="个", use_unit="测试三矿",
       req_date=dt.date(2026, 11, 25), remark="新需求，等价格", qty=6),
    tr("TEST-2002", "测试喷头", spec="雾化喷头 不锈钢", unit="只", use_unit="测试三矿",
       req_date=dt.date(2026, 11, 28), qty=50),
])
wb2.save(OUT / "测试文件2_仅需求无报价.xlsx")

# ============ 文件3：手工乱表（列位移 + 横向多供应商 + 序列号日期） ============
wb3 = Workbook()
ws3 = wb3.active
ws3.title = "Sheet1"
# 表头不含"行号"、数据行首列是行号 → 触发位移检测；尾部 小红/小明 为横向供应商组
ws3.append(["物料编码", "物料描述", "规格型号", "技术参数", "品牌", "单位", "零件号/图号",
            "材质", "进口/国产", "数量", "使用单位", "需求日期", "备注", "", "小红", "", "小明"])
# 数据：[行号, 编码, 描述, 规格, 品牌, 单位, 零件, 材质, 进口国产, 数量, 使用单位, 需求日期(序列号),
#        备注, "", 电话, 价格&店名…, 小明组]
ws3.append([1, "TEST-3001", "测试皮带轮||传动件|SPB355-3|3535", "", "", "件", "", "", "",
            8, "测试四矿", 46357, "张工", "", "13812345678",
            "微信单价320含税运 阳光机械器材销售处", "", "350 含税运 宏达五金店"])
ws3.append([2, "TEST-3002", "测试密封件||密封件|O型圈|丁晴橡胶", "O型圈 GB3452.1", "", "件",
            "", "", "", 100, "测试四矿", 46357, "",
            "", "", "每件2.5元 不含运 顺发橡塑制品厂", "", "2.3 含税运 日升密封件经营部"])
ws3.append([3, "TEST-3003", "测试联轴器||传动件|梅花联轴器", "LM-3", "", "个", "", "", "",
            2, "测试四矿", 46358, "等你报价"])
ws3.append([4, "TEST-3004", "测试钢管||钢管件|DN100|20号钢", "DN100x4", "", "米", "", "", "",
            50, "测试五矿", 46358])
wb3.save(OUT / "测试文件3_手工乱表_横向拆分.xlsx")

# ============ 文件4：异常与特殊行 ============
wb4 = Workbook()
ws4 = wb4.active
ws4.title = "Sheet1"
ws4.append(["行号", "物料编码", "物料描述", "规格型号", "技术参数", "品牌", "单位", "零件号/图号",
            "材质", "进口/国产", "数量", "接收人", "使用单位", "需求日期", "备注", "定价"])
ws4.append([1, "TEST-4001", "测试齿轮||传动件|模数4|45号钢", "m4 z58", "", "", "件", "", "", "",
            3, "李工", "测试六矿", "2026-12-05", "正常备注 待报价：8300 含税运", 8100])
ws4.append([2, "", "", "", "", "", "", "", "", "", "", "", "", "", "名称编码都为空的异常行"])  # 错误行
ws4.append([3, "TEST-4003", "测试垫片||密封件|石棉", "JBT-98", "", "", "件", "", "", "",
            10, "王工", "测试六矿", "2026-12-06", "厂家没人回"])  # 无有效报价
ws4.append([4, "TEST-4004", "测试弹簧||弹性件|65Mn", "5x30", "", "", "件", "", "", "",
            5, "王工", "测试六矿", "2026-12-06", "电话13957934455 8300 含税运", ""])  # 备注提取
ws4.append([5, "", "", "", "", "", "", "", "", "", "", "", "", "", ""])  # 错误行
ws4.append([6, "TEST-4006", "测试卡箍||紧固件|不锈钢", "抱箍50", "", "", "只", "", "", "",
            30, "赵工", "测试六矿", 46360, "电话1337564321 单价9块，含税", ""])  # 序列号日期 + 备注提取
wb4.save(OUT / "测试文件4_异常与特殊行.xlsx")

print("生成完成：")
for f in sorted(OUT.glob("*.xlsx")):
    print(" -", f.name)
